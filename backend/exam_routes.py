import base64
import hashlib
import io
import json
import random
import secrets
import socket
import sqlite3
import jwt
import qrcode
import docx_import
import python as P  

from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from math import ceil
from typing import Literal
from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

Difficulty = Literal["NB", "TH", "VD", "VDC"]
QType = Literal["multiple_choice", "true_false", "short_answer"]
FMT = "%Y-%m-%d %H:%M:%S"  

router = APIRouter(prefix="/api")
teacher = APIRouter(dependencies=[Depends(P.get_current_user)])

@contextmanager
def db():
    conn = P.get_db()
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()

def require_resource_owner(request: Request, user=Depends(P.get_current_user)):
    if user["role"] == "admin":
        return user

    params = request.path_params
    with db() as c:
        if "exam_id" in params:
            row = c.execute(
                "SELECT created_by FROM exams WHERE exam_id=?",
                (params["exam_id"],),
            ).fetchone()
        elif "sid" in params:
            row = c.execute(
                "SELECT created_by FROM exam_sessions WHERE session_id=?",
                (params["sid"],),
            ).fetchone()
            if row and "ssid" in params:
                student = c.execute(
                    "SELECT 1 FROM session_students WHERE session_student_id=? AND session_id=?",
                    (params["ssid"], params["sid"]),
                ).fetchone()
                if not student:
                    raise HTTPException(404, "Không tìm thấy thí sinh trong phòng thi")
        else:
            return user

    if not row:
        raise HTTPException(404, "Không tìm thấy tài nguyên")
    if row["created_by"] != user["user_id"]:
        raise HTTPException(404, "Không tìm thấy tài nguyên")
    return user

teacher.dependencies.append(Depends(require_resource_owner))

def now_utc():
    return datetime.now(timezone.utc)

def parse(s):
    return datetime.strptime(s, FMT).replace(tzinfo=timezone.utc)

def deadline(started_at, minutes, end_time):
    d = parse(started_at) + timedelta(minutes=minutes)
    return min(d, parse(end_time)) if end_time else d

def sha(
    *parts,
):  # xáo trộn ổn định theo (học sinh, câu hỏi): F5 hay vào lại vẫn ra cùng thứ tự
    return hashlib.sha256(":".join(map(str, parts)).encode()).hexdigest()

def lan_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(
            ("10.255.255.255", 1)
        )  # không gửi gói nào, chỉ để OS chọn card mạng LAN
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()

def validate_question_data(data):
    if not data.question_text.strip():
        raise HTTPException(
            status_code=400, detail="nội dung câu hỏi không được để trống"
        )

    t, opts = data.question_type, data.options

    if len({o.option_key for o in opts}) != len(opts) or len(
        {o.sort_order for o in opts}
    ) != len(opts):
        raise HTTPException(
            status_code=400, detail="option_key và sort_order không được trùng nhau"
        )

    if t == "short_answer":
        if opts:
            raise HTTPException(
                status_code=400, detail="câu hỏi trả lời ngắn không dùng options"
            )
        if not data.accepted_answers:
            raise HTTPException(
                status_code=400,
                detail="câu hỏi trả lời ngắn phải có ít nhất 1 đáp án được chấp nhận",
            )
        return
    if data.accepted_answers:
        raise HTTPException(
            status_code=400, detail="loại câu hỏi này không dùng accepted_answers"
        )
    if t == "multiple_choice":
        if len(opts) < 2 or sum(o.is_correct for o in opts) != 1:
            raise HTTPException(
                status_code=400,
                detail="câu trắc nghiệm cần ≥2 lựa chọn và đúng 1 đáp án đúng",
            )
    elif {o.option_key for o in opts} != {"A", "B", "C", "D"}:
        raise HTTPException(
            status_code=400,
            detail="câu đúng/sai phải có đúng 4 ý với option_key A, B, C, D",
        )


# phần lớp và học sinh
class ClassCreate(BaseModel):
    class_name: str = Field(min_length=1, max_length=50)
    grade_id: int
    school_year: str | None = Field(default=None, max_length=20)


class StudentCreate(BaseModel):
    student_code: str = Field(min_length=1, max_length=50)
    full_name: str = Field(min_length=1, max_length=100)
    date_of_birth: date | None = None
    class_id: int


def insert_or_409(c, sql, args, msg):
    try:
        return c.execute(sql, args).lastrowid
    except sqlite3.IntegrityError:
        raise HTTPException(409, msg)


@teacher.post("/classes", status_code=201)
def create_class(d: ClassCreate, user=Depends(P.get_current_user)):
    with db() as c:
        cid = insert_or_409(
            c,
            """ 
        INSERT INTO classes (
        class_name, 
        grade_id, 
        school_year,
        created_by)
        VALUES (?,?,?,?)
        """,
            (d.class_name.strip(), d.grade_id, d.school_year, user["user_id"]),
            "Lớp đã tồn tại hoặc khối lớp không hợp lệ",
        )
        return {"class_id": cid}


@teacher.post("/students", status_code=201)
def create_student(d: StudentCreate, user=Depends(P.get_current_user)):
    with db() as c:
        owned_class = c.execute(
            "SELECT 1 FROM classes WHERE class_id=? AND (?='admin' OR created_by=?)",
            (d.class_id, user["role"], user["user_id"]),
        ).fetchone()
        if not owned_class:
            raise HTTPException(404, "Không tìm thấy lớp học")
        sid = insert_or_409(
            c,
            """
        INSERT INTO students (
        student_code, 
        full_name, 
        date_of_birth,
        class_id) 
        VALUES (?,?,?,?)
        """,
            (
                d.student_code.strip(),
                d.full_name.strip(),
                d.date_of_birth.isoformat() if d.date_of_birth else None,
                d.class_id,
            ),
            "Mã học sinh đã tồn tại hoặc lớp không hợp lệ",
        )
        return {"student_id": sid}


# phần tạo đề theo ma trận
class MatrixRow(BaseModel):
    topic_id: int
    lesson_id: int | None = None
    difficulty: Difficulty
    question_type: QType
    question_count: int = Field(gt=0, le=200)


class ExamCreate(BaseModel):
    title: str = Field(min_length=1, max_length=255)
    supervising_authority: str = Field(default="", max_length=255)
    school_name: str = Field(default="", max_length=255)
    subject_id: int
    grade_id: int
    duration_minutes: int = Field(gt=0)
    matrix: list[MatrixRow] = Field(min_length=1, max_length=50)


def get_exam(c, exam_id, draft=False):
    e = c.execute("SELECT * FROM exams WHERE exam_id=?", (exam_id,)).fetchone()
    if not e:
        raise HTTPException(404, "Không tìm thấy đề thi")
    if draft and e["status"] != "draft":
        raise HTTPException(409, "Đề đã lưu, không sửa được nữa")
    return e


def candidates(c, e, topic_id, lesson_id, difficulty, qtype):
    sql = """
        SELECT 
        question_id 
        FROM questions 
        WHERE is_active=1 
            AND subject_id=? 
            AND grade_id=? 
            AND topic_id=? 
            AND difficulty=? 
            AND question_type=?
            AND (created_by=? OR is_public=1)
        """

    args = [e["subject_id"], e["grade_id"], topic_id, difficulty, qtype, e["created_by"]]

    if lesson_id:
        sql += " AND lesson_id=?"
        args.append(lesson_id)
    return [r[0] for r in c.execute(sql, args)]


def base_version(c, exam_id):  # "đề gốc" = version có mã GOC
    r = c.execute(
        """
        SELECT 
        version_id 
        FROM exam_versions 
        WHERE exam_id=? 
            AND version_code='GOC'
        """,
        (exam_id,),
    ).fetchone()

    if r:
        count = c.execute(
            "SELECT COUNT(*) FROM exam_version_questions WHERE version_id=?", (r[0],)
        ).fetchone()[0]
        legacy = c.execute(
            """
            SELECT version_id FROM exam_versions
            WHERE exam_id=? AND version_code='A'
            """,
            (exam_id,),
        ).fetchone()
        if not count and legacy:
            c.execute(
                """
                INSERT OR IGNORE INTO exam_version_questions
                    (version_id, question_id, question_order, score)
                SELECT ?, question_id, question_order, score
                FROM exam_version_questions WHERE version_id=?
                """,
                (r[0], legacy[0]),
            )
        return r[0]

    legacy = c.execute(
        "SELECT version_id FROM exam_versions WHERE exam_id=? AND version_code='A'",
        (exam_id,),
    ).fetchone()
    if legacy:
        c.execute(
            "UPDATE exam_versions SET version_code='GOC' WHERE version_id=?",
            (legacy[0],),
        )
        return legacy[0]

    return c.execute(
        """
        INSERT INTO exam_versions (
        exam_id, 
        version_code
        ) VALUES (?, 'GOC')
        """,
        (exam_id,),
    ).lastrowid


def write_base(c, exam_id, ids):
    vid = base_version(c, exam_id)
    c.execute(
        """
        DELETE 
        FROM exam_version_questions 
        WHERE version_id=?
        """,
        (vid,),
    )
    # shortcut: chia đều 10 điểm cho mọi câu, chưa có trọng số theo loại câu; nâng cấp khi có quy định điểm
    score = round(10 / len(ids), 4)
    c.executemany(
        """
        INSERT INTO exam_version_questions (
        version_id, 
        question_id, 
        question_order, 
        score
        ) 
        VALUES (?,?,?,?)
        """,
        [(vid, q, i + 1, score) for i, q in enumerate(ids)],
    )


def preview(c, exam_id):
    e = get_exam(c, exam_id)
    qs = c.execute(
        """
        SELECT 
        evq.question_order, 
        evq.score, q.question_id, 
        q.question_type, 
        q.difficulty, 
        q.topic_id, 
        q.lesson_id, 
        q.question_text
        FROM exam_versions v
        JOIN exam_version_questions evq ON evq.version_id = v.version_id
        JOIN questions q ON q.question_id = evq.question_id
        WHERE v.exam_id=? 
            AND v.version_code='GOC' 
        ORDER BY evq.question_order
        """,
        (exam_id,),
    ).fetchall()

    matrix = c.execute(
        """
        SELECT * 
        FROM exam_matrices 
        WHERE exam_id=? 
        ORDER BY matrix_id
        """,
        (exam_id,),
    ).fetchall()

    return {
        "exam": dict(e),
        "matrix": [dict(m) for m in matrix],
        "questions": [dict(q) for q in qs],
    }


@teacher.post("/exams", status_code=201)
def create_exam(d: ExamCreate, user: dict = Depends(P.get_current_user)):
    with db() as c:
        try:
            eid = c.execute(
                """
                INSERT INTO exams (
                title, 
                supervising_authority,
                school_name,
                subject_id, 
                grade_id, 
                duration_minutes, 
                created_by
                ) 
                VALUES (?,?,?,?,?,?,?)
                """,
                (
                    d.title.strip(),
                    d.supervising_authority.strip(),
                    d.school_name.strip(),
                    d.subject_id,
                    d.grade_id,
                    d.duration_minutes,
                    user["user_id"],
                ),
            ).lastrowid

            c.executemany(
                """
                INSERT INTO exam_matrices (
                exam_id, 
                topic_id, 
                lesson_id, 
                difficulty, 
                question_type, 
                question_count
                ) 
                VALUES (?,?,?,?,?,?)
                """,
                [
                    (
                        eid,
                        m.topic_id,
                        m.lesson_id,
                        m.difficulty,
                        m.question_type,
                        m.question_count,
                    )
                    for m in d.matrix
                ],
            )

        except sqlite3.IntegrityError:
            raise HTTPException(
                status_code=400,
                detail="Môn/khối/chủ đề/bài học không hợp lệ hoặc tiêu đề trống",
            )

        return {"exam_id": eid}


@teacher.get("/exams")
def list_exams(user=Depends(P.get_current_user)):
    with db() as c:
        if user["role"] == "admin":
            rows = c.execute("SELECT * FROM exams ORDER BY exam_id DESC")
        else:
            rows = c.execute(
                "SELECT * FROM exams WHERE created_by=? ORDER BY exam_id DESC",
                (user["user_id"],),
            )
        return {"exams": [dict(r) for r in rows]}


@teacher.get("/exams/{exam_id}")
def exam_detail(exam_id: int):
    with db() as c:
        return preview(c, exam_id)


@teacher.post("/exams/{exam_id}/generate")
def generate_exam(exam_id: int):
    """Rút ngẫu nhiên câu hỏi theo ma trận; gọi lại để làm lại toàn bộ đề gốc."""
    with db() as c:
        e = get_exam(c, exam_id, draft=True)
        picked = []
        for m in c.execute(
            """
            SELECT * 
            FROM exam_matrices 
            WHERE exam_id=? 
            ORDER BY matrix_id
            """,
            (exam_id,),
        ).fetchall():

            pool = [
                q
                for q in candidates(
                    c,
                    e,
                    m["topic_id"],
                    m["lesson_id"],
                    m["difficulty"],
                    m["question_type"],
                )
                if q not in picked
            ]
            if len(pool) < m["question_count"]:
                raise HTTPException(
                    400,
                    f"Thiếu câu hỏi (chủ đề {m['topic_id']}, {m['difficulty']}, {m['question_type']}): "
                    f"cần {m['question_count']}, ngân hàng chỉ có {len(pool)}",
                )

            picked += random.sample(pool, m["question_count"])
        write_base(c, exam_id, picked)
        return preview(c, exam_id)


class Swap(BaseModel):
    question_id: int


@teacher.post("/exams/{exam_id}/swap")
def swap_question(exam_id: int, d: Swap):
    """Đổi 1 câu trong đề gốc sang câu khác cùng dòng ma trận."""
    with db() as c:
        e = get_exam(c, exam_id, draft=True)
        vid = base_version(c, exam_id)
        cur = [
            r[0]
            for r in c.execute(
                """
            SELECT 
            question_id 
            FROM exam_version_questions 
            WHERE version_id=?
            """,
                (vid,),
            )
        ]

        q = c.execute(
            """
            SELECT 
            topic_id, 
            lesson_id, 
            difficulty, 
            question_type 
            FROM questions 
            WHERE question_id=?
            """,
            (d.question_id,),
        ).fetchone()

        if d.question_id not in cur or not q:
            raise HTTPException(status_code=404, detail="Câu hỏi không có trong đề gốc")

        m = c.execute(
            """
            SELECT 
            lesson_id 
            FROM exam_matrices 
            WHERE exam_id=? 
                AND topic_id=? 
                AND difficulty=? 
                AND question_type=? 
            """
            """
            AND (lesson_id IS NULL OR lesson_id=?)
            """,
            (
                exam_id,
                q["topic_id"],
                q["difficulty"],
                q["question_type"],
                q["lesson_id"],
            ),
        ).fetchone()

        if not m:
            raise HTTPException(
                status_code=400, detail="Không tìm thấy dòng ma trận của câu này"
            )

        pool = [
            x
            for x in candidates(
                c, e, q["topic_id"], m["lesson_id"], q["difficulty"], q["question_type"]
            )
            if x not in cur
        ]

        if not pool:
            raise HTTPException(
                status_code=409, detail="Ngân hàng không còn câu nào khác để thay"
            )

        c.execute(
            """
            UPDATE exam_version_questions 
            SET question_id=? 
            WHERE version_id=? 
                AND question_id=?""",
            (random.choice(pool), vid, d.question_id),
        )

        return preview(c, exam_id)


@teacher.post("/exams/{exam_id}/save")
def save_exam(exam_id: int):
    with db() as c:
        get_exam(c, exam_id, draft=True)
        n = c.execute(
            """SELECT COUNT(*) 
            FROM exam_version_questions evq 
            JOIN exam_versions v 
            ON v.version_id=evq.version_id """
            """
            WHERE v.exam_id=? 
                AND v.version_code='GOC'
            """,
            (exam_id,),
        ).fetchone()[0]

        if n == 0:
            raise HTTPException(status_code=400, detail="Chưa sinh đề gốc")

        c.execute(
            """
            UPDATE exams 
            SET status='ready', 
                updated_at=? 
            WHERE exam_id=?
            """,
            (P.get_utc_time(), exam_id),
        )

        return {"status": "ok", "exam_id": exam_id, "question_count": n}


def exam_question_details(c, exam_id):
    data = preview(c, exam_id)
    if not data["questions"]:
        raise HTTPException(400, "Chưa sinh đề gốc")
    names = c.execute(
        """
        SELECT s.subject_name, g.grade_name
        FROM exams e
        JOIN subjects s ON s.subject_id = e.subject_id
        JOIN grades g ON g.grade_id = e.grade_id
        WHERE e.exam_id=?
        """,
        (exam_id,),
    ).fetchone()
    data["subject_name"] = names["subject_name"] if names else ""
    data["grade_name"] = names["grade_name"] if names else ""
    for q in data["questions"]:
        q["options"] = [
            dict(o)
            for o in c.execute(
                """
            SELECT option_key, option_text, is_correct
            FROM question_options
            WHERE question_id=?
            ORDER BY sort_order
            """,
                (q["question_id"],),
            )
        ]
        q["accepted_answers"] = [
            r[0]
            for r in c.execute(
                """
            SELECT answer_text
            FROM question_accepted_answers
            WHERE question_id=?
            """,
                (q["question_id"],),
            )
        ]
    return data


def exam_answer_line(q):
    if q["question_type"] == "multiple_choice":
        keys = [o["option_key"] for o in q["options"] if o["is_correct"]]
        return keys[0] if keys else ""
    if q["question_type"] == "true_false":
        return " ".join(
            f'{o["option_key"]}={"Đ" if o["is_correct"] else "S"}' for o in q["options"]
        )
    return "; ".join(q["accepted_answers"])


def exam_paper_docx(data):
    e = data["exam"]
    blocks = [
        {
            "text": (e.get("supervising_authority") or "").upper(),
            "align": "center",
            "bold": True,
            "size": 24,
        },
        {
            "text": (e.get("school_name") or "").upper(),
            "align": "center",
            "bold": True,
            "size": 24,
        },
        {"text": "", "align": "center"},
        {"text": "ĐỀ THI GỐC", "align": "center", "bold": True, "size": 32},
        {"text": e["title"], "align": "center", "bold": True, "size": 28},
        {
            "text": f'{data.get("subject_name") or ""} · {data.get("grade_name") or ""} · Thời gian: {e["duration_minutes"]} phút',
            "align": "center",
            "size": 22,
        },
        {
            "text": f'Mã đề: GOC · {len(data["questions"])} câu',
            "align": "center",
            "size": 22,
        },
        "",
    ]
    for q in data["questions"]:
        blocks.append(
            {
                "text": f'Câu {q["question_order"]}. {q["question_text"]}',
                "bold": True,
                "size": 22,
            }
        )
        if q["question_type"] == "true_false":
            for o in q["options"]:
                blocks.append(f'    {o["option_key"]}) {o["option_text"]}')
            blocks.append("    Đánh Đ/S cho từng ý.")
        else:
            for o in q["options"]:
                blocks.append(f'    {o["option_key"]}. {o["option_text"]}')
        blocks.append("")
    blocks.append(
        {"text": "ĐÁP ÁN ĐỀ GỐC", "align": "center", "bold": True, "size": 28}
    )
    for q in data["questions"]:
        blocks.append(f'Câu {q["question_order"]}: {exam_answer_line(q)}')
    return docx_import.paragraphs_docx(blocks)


@teacher.get("/exams/{exam_id}/export.docx")
def export_exam(exam_id: int):
    with db() as c:
        data = exam_question_details(c, exam_id)
        return Response(
            exam_paper_docx(data),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={
                "Content-Disposition": f'attachment; filename="de_goc_{exam_id}.docx"'
            },
        )


# phần phòng thi
class SessionCreate(BaseModel):
    exam_id: int
    session_name: str = Field(min_length=1, max_length=100)
    start_time: datetime | None = None  # giờ không kèm múi giờ được hiểu là giờ máy chủ
    end_time: datetime | None = None
    password: str | None = Field(default=None, min_length=8, max_length=128)
    shuffle_questions: bool = True
    shuffle_options: bool = True
    paper_code_prefix: str = Field(default="", max_length=8)
    paper_code_length: int = Field(default=3, ge=1, le=8)
    paper_variant_count: int = Field(default=1, ge=1, le=100)


def to_db(dt):
    return dt.astimezone(timezone.utc).strftime(FMT) if dt else None


def get_session(c, sid):
    s = c.execute(
        """
        SELECT * 
        FROM exam_sessions 
        WHERE session_id=?
        """,
        (sid,),
    ).fetchone()

    if not s:
        raise HTTPException(status_code=404, detail="Không tìm thấy phòng thi")
    return s


@teacher.post("/sessions", status_code=201)
def create_session(
    d: SessionCreate, request: Request, user: dict = Depends(P.get_current_user)
):
    if d.start_time and d.end_time and d.end_time <= d.start_time:
        raise HTTPException(status_code=400, detail="Giờ kết thúc phải sau giờ bắt đầu")
    prefix = d.paper_code_prefix.strip().upper()
    if prefix and (not prefix.isascii() or not prefix.isalnum()):
        raise HTTPException(400, "Tiền tố mã đề chỉ gồm chữ cái và chữ số ASCII")
    if d.paper_variant_count > 10**d.paper_code_length:
        raise HTTPException(400, "Độ dài mã không đủ chứa số đề muốn trộn")
    if d.paper_variant_count > 1 and not (d.shuffle_questions or d.shuffle_options):
        raise HTTPException(
            400, "Bật trộn câu hỏi hoặc phương án để tạo nhiều mã đề khác nhau"
        )

    with db() as c:
        e = get_exam(c, d.exam_id)
        if user["role"] != "admin" and e["created_by"] != user["user_id"]:
            raise HTTPException(404, "Không tìm thấy đề thi")
        if e["status"] != "ready":
            raise HTTPException(
                status_code=409, detail="Đề chưa được lưu (status phải là ready)"
            )

        base = base_version(c, d.exam_id)
        if not c.execute(
            "SELECT 1 FROM exam_version_questions WHERE version_id=? LIMIT 1", (base,)
        ).fetchone():
            raise HTTPException(
                409, "Đề gốc chưa có câu hỏi. Hãy sinh và lưu đề trước khi tạo phòng."
            )
        # mỗi phòng có 1 bản chụp đề riêng: sửa/đổi đề gốc sau này không ảnh hưởng phòng đang thi
        vid = c.execute(
            """
            INSERT INTO exam_versions (
            exam_id, 
            version_code, 
            shuffle_questions, 
            shuffle_options
            ) 
            VALUES (?,?,?,?)
            """,
            (
                d.exam_id,
                "S" + secrets.token_hex(3).upper(),
                d.shuffle_questions,
                d.shuffle_options,
            ),
        ).lastrowid

        c.execute(
            """
            INSERT INTO exam_version_questions (
            version_id, 
            question_id, 
            question_order, 
            score) 
            """
            """
            SELECT 
            ?, 
            question_id, 
            question_order,
            score 
            FROM exam_version_questions 
            WHERE version_id=?
            """,
            (vid, base),
        )

        code = secrets.token_hex(8).upper()
        sid = c.execute(
            """
            INSERT INTO exam_sessions (
            exam_id, 
            session_name, 
            version_id, 
            paper_code_prefix,
            paper_code_length,
            paper_variant_count,
            access_code, 
            password_hash, 
            start_time, 
            end_time, 
            created_by
            )
            VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                d.exam_id,
                d.session_name.strip(),
                vid,
                prefix,
                d.paper_code_length,
                d.paper_variant_count,
                code,
                P.password_hash.hash(d.password) if d.password else None,
                to_db(d.start_time),
                to_db(d.end_time),
                user["user_id"],
            ),
        ).lastrowid

        return {
            "session_id": sid,
            "access_code": code,
            "join_url": join_url(request, code),
            "qr_data_url": join_qr_data_url(request, code),
        }


def join_url(request, code):
    return (
        f"{request.url.scheme}://{lan_ip()}:{request.url.port or 80}/join?code={code}"
    )


def join_qr_data_url(request, code):
    image = qrcode.make(join_url(request, code))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return "data:image/png;base64," + base64.b64encode(output.getvalue()).decode()


@teacher.get("/sessions")
def list_sessions(user=Depends(P.get_current_user)):
    with db() as c:
        sql = """
            SELECT session_id, exam_id, session_name, access_code,
                   start_time, end_time, status
            FROM exam_sessions
        """
        args = ()
        if user["role"] != "admin":
            sql += " WHERE created_by=?"
            args = (user["user_id"],)
        sql += " ORDER BY session_id DESC"
        rows = c.execute(sql, args)
        return {"sessions": [dict(r) for r in rows]}


@teacher.get("/sessions/{sid}")
def session_detail(sid: int, request: Request):
    with db() as c:
        s = dict(get_session(c, sid))
        s.pop("password_hash")
        s["has_password"] = bool(get_session(c, sid)["password_hash"])
        s["join_url"] = join_url(request, s["access_code"])
        s["qr_data_url"] = join_qr_data_url(request, s["access_code"])
        return s


class SessionPassword(BaseModel):
    password: str = Field(min_length=8, max_length=128)


@teacher.put("/sessions/{sid}/password")
def set_session_password(sid: int, d: SessionPassword):
    with db() as c:
        get_session(c, sid)
        c.execute(
            "UPDATE exam_sessions SET password_hash=? WHERE session_id=?",
            (P.password_hash.hash(d.password), sid),
        )
        c.execute(
            "DELETE FROM student_room_login_attempts WHERE session_student_id IN "
            "(SELECT session_student_id FROM session_students WHERE session_id=?)",
            (sid,),
        )
        return {"status": "ok"}


class StatusChange(BaseModel):
    status: Literal["ongoing", "completed", "cancelled"]


NEXT = {"scheduled": {"ongoing", "cancelled"}, "ongoing": {"completed", "cancelled"}}


@teacher.post("/sessions/{sid}/status")
def change_status(sid: int, d: StatusChange):
    with db() as c:
        s = get_session(c, sid)
        if d.status not in NEXT.get(s["status"], set()):
            raise HTTPException(
                status_code=409,
                detail=f"Không thể chuyển từ {s['status']} sang {d.status}",
            )

        if d.status == "completed":  # thu bài những ai còn đang làm, ai chưa vào = vắng
            for (ssid,) in c.execute(
                """
                SELECT 
                session_student_id 
                FROM session_students 
                WHERE session_id=?
                    AND status='in_progress'
                """,
                (sid,),
            ).fetchall():

                finalize(c, ssid)

            c.execute(
                """
                UPDATE session_students 
                SET status='absent' 
                WHERE session_id=? 
                    AND status='registered'
                """,
                (sid,),
            )

        if d.status == "ongoing" and not s["start_time"]:
            c.execute(
                "UPDATE exam_sessions SET start_time=? WHERE session_id=?",
                (now_utc().strftime(FMT), sid),
            )

        c.execute(
            """
            UPDATE exam_sessions 
            SET status=? 
            WHERE session_id=?
            """,
            (d.status, sid),
        )

        return {"status": "ok", "session_status": d.status}


# phần danh sách thí sinh 
def add_to_session(c, sid, student_id):
    c.execute(
        """
        INSERT OR IGNORE INTO session_students (
        session_id, 
        student_id, 
        student_code_snapshot, 
        name_snapshot, 
        class_snapshot
        )
        SELECT ?, 
        s.student_id, 
        s.student_code, 
        s.full_name, 
        cl.class_name
        FROM students s 
        JOIN classes cl ON cl.class_id = s.class_id 
        WHERE s.student_id=?
        """,
        (sid, student_id),
    )


class AddStudents(BaseModel):
    student_ids: list[int] = Field(default_factory=list, max_length=1000)
    class_id: int | None = None


def open_session(c, sid):
    s = get_session(c, sid)
    if s["status"] not in ("scheduled", "ongoing"):
        raise HTTPException(status_code=409, detail="Phòng thi đã kết thúc")
    return s


@teacher.post("/sessions/{sid}/students")
def add_students(sid: int, d: AddStudents, user=Depends(P.get_current_user)):
    with db() as c:
        open_session(c, sid)
        ids = list(dict.fromkeys(d.student_ids))
        if d.class_id:
            owned_class = c.execute(
                "SELECT 1 FROM classes WHERE class_id=? AND (?='admin' OR created_by=?)",
                (d.class_id, user["role"], user["user_id"]),
            ).fetchone()
            if not owned_class:
                raise HTTPException(404, "Không tìm thấy lớp học")
            ids += [
                r[0]
                for r in c.execute(
                    """
                SELECT student_id 
                FROM students 
                WHERE class_id=? 
                    AND is_active=1
                """,
                    (d.class_id,),
                )
            ]
        ids = list(dict.fromkeys(ids))
        if ids:
            marks = ",".join("?" for _ in ids)
            valid_ids = {
                row[0]
                for row in c.execute(
                    f"SELECT s.student_id FROM students s JOIN classes cl ON cl.class_id=s.class_id "
                    f"WHERE s.student_id IN ({marks}) AND (?='admin' OR cl.created_by=?)",
                    (*ids, user["role"], user["user_id"]),
                )
            }
            if valid_ids != set(ids):
                raise HTTPException(404, "Không tìm thấy học sinh trong lớp của bạn")
        for i in ids:
            add_to_session(c, sid, i)

        return {"status": "ok", "requested": len(ids)}


@teacher.delete("/sessions/{sid}/students/{ssid}")
def remove_student(sid: int, ssid: int):
    with db() as c:
        cur = c.execute(
            """
            DELETE FROM session_students 
            WHERE session_id=? 
                AND session_student_id=? 
                AND status='registered'
            """,
            (sid, ssid),
        )

        if not cur.rowcount:
            raise HTTPException(
                status_code=409,
                detail="Không tìm thấy hoặc học sinh đã vào thi, không xóa được",
            )

        return {"status": "ok"}


def cell(x):
    if isinstance(x, float) and x.is_integer():
        x = int(x)
    return "" if x is None else str(x).strip()


def parse_student_date(value):
    value = cell(value)
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Ngày sinh phải có dạng DD/MM/YYYY hoặc YYYY-MM-DD")


def parse_roster(c, raw, user):
    try:
        table = docx_import.table_rows(raw)
    except Exception as e:
        raise HTTPException(
            400, "File không phải Word (.docx) có bảng danh sách hợp lệ"
        ) from e
    expected = ("stt", "mã học sinh / sbd", "họ và tên", "lớp")
    header = tuple(" ".join(x.split()).casefold() for x in (table[0] if table else []))
    if header[:4] != expected:
        raise HTTPException(
            400,
            "Bảng Word cần có các cột theo thứ tự: STT | Mã học sinh / SBD | Họ và tên | Lớp | Ngày sinh",
        )
    if len(table) > 1001:
        raise HTTPException(413, "Danh sách tối đa 1.000 học sinh")

    rows, seen = [], set()
    for i, r in enumerate(table[1:], start=2):
        code, name, cls, dob_text = (cell(x) for x in (list(r) + [None] * 5)[1:5])
        if not (code or name or cls or dob_text):
            continue
        class_id, err, dob = None, None, None
        if not (code and name and cls):
            err = "Thiếu thông tin bắt buộc"
        elif code in seen:
            err = "Trùng mã học sinh/SBD trong file"
        else:
            seen.add(code)
            try:
                dob = parse_student_date(dob_text)
            except ValueError as e:
                err = str(e)
            if not err:
                k = c.execute(
                    "SELECT class_id FROM classes WHERE class_name=? AND is_active=1 AND (?='admin' OR created_by=?)",
                    (cls, user["role"], user["user_id"]),
                ).fetchone()
                class_id = k[0] if k else None
                err = None if k else f"Lớp '{cls}' chưa tồn tại"
        rows.append(
            {
                "row": i,
                "student_code": code,
                "full_name": name,
                "date_of_birth": dob,
                "class_name": cls,
                "class_id": class_id,
                "error": err,
            }
        )
    return rows

@teacher.post("/sessions/{sid}/students/import")
def import_students(
    sid: int,
    file: UploadFile = File(...),
    commit: bool = False,
    user=Depends(P.get_current_user),
):
    """commit=false: chỉ xem trước + báo lỗi. commit=true: ghi vào phòng thi, chỉ khi file không còn lỗi."""
    if not (file.filename or "").lower().endswith(".docx"):
        raise HTTPException(400, "Chỉ nhận file Word (.docx)")
    with db() as c:
        open_session(c, sid)
        raw = P.read_upload_limited(file, 5 * 1024 * 1024, "Danh sách tối đa 5MB")
        rows = parse_roster(c, raw, user)
        errors = [r for r in rows if r["error"]]
        if not commit:
            return {"count": len(rows), "error_count": len(errors), "rows": rows}
        if errors or not rows:
            raise HTTPException(
                400,
                {
                    "message": "File còn lỗi hoặc rỗng, hãy xem trước và sửa lại",
                    "errors": errors,
                },
            )
        for r in rows:
            st = c.execute(
                """
                SELECT s.student_id, s.date_of_birth
                FROM students s JOIN classes cl ON cl.class_id=s.class_id
                WHERE s.student_code=? AND (?='admin' OR cl.created_by=?)
                """,
                (r["student_code"], user["role"], user["user_id"]),
            ).fetchone()

            if not st and c.execute(
                "SELECT 1 FROM students WHERE student_code=?",
                (r["student_code"],),
            ).fetchone():
                raise HTTPException(409, "Mã học sinh đã được sử dụng")

            if st:
                student_id = st["student_id"]
                if r["date_of_birth"]:
                    c.execute(
                        "UPDATE students SET date_of_birth=? WHERE student_id=?",
                        (r["date_of_birth"], student_id),
                    )
            else:
                student_id = c.execute(
                    """
                    INSERT INTO students (student_code, full_name, date_of_birth, class_id)
                    VALUES (?,?,?,?)
                    """,
                    (
                        r["student_code"],
                        r["full_name"],
                        r["date_of_birth"],
                        r["class_id"],
                    ),
                ).lastrowid
            add_to_session(c, sid, student_id)

        return {"status": "ok", "imported": len(rows)}

@router.get("/templates/students.docx")
def students_template():
    data = docx_import.table_docx(
        ["STT", "Mã học sinh / SBD", "Họ và tên", "Lớp", "Ngày sinh"],
        ["1", "", "", "", ""],
    )
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": 'attachment; filename="danh_sach_hoc_sinh_mau.docx"'
        },
    )

def xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

@teacher.get("/sessions/{sid}/monitor")
def monitor(sid: int):
    with db() as c:
        get_session(c, sid)
        ensure_graded(c, sid)
        rows = c.execute(
            """
            SELECT 
            ss.session_student_id, 
            ss.student_code_snapshot AS student_code, 
            ss.name_snapshot AS full_name,
            ss.class_snapshot AS class_name, 
            ss.status, 
            ss.started_at, 
            ss.submitted_at,
            r.total_score, 
            r.correct_count,

            (
            SELECT COUNT(*) 
            FROM answers a 
            WHERE a.session_student_id = ss.session_student_id
            ) AS answered,

            (
            SELECT COUNT(*) 
            FROM exam_violations v 
            WHERE v.session_student_id = ss.session_student_id
            ) AS violations

            FROM session_students ss LEFT JOIN results r ON r.session_student_id = ss.session_student_id
            WHERE ss.session_id=? 
            ORDER BY ss.name_snapshot
            """,
            (sid,),
        ).fetchall()

        count = lambda st: sum(r["status"] == st for r in rows)

        return {
            "total": len(rows),
            "not_joined": count("registered") + count("absent"),
            "in_progress": count("in_progress"),
            "submitted": count("submitted"),
            "students": [dict(r) for r in rows],
            "absent_names": [
                r["full_name"] for r in rows if r["status"] in ("registered", "absent")
            ],
        }


@teacher.get("/sessions/{sid}/students")
def session_students(sid: int):
    with db() as c:
        get_session(c, sid)
        rows = c.execute(
            """
            SELECT student_code_snapshot AS student_code,
                   name_snapshot AS full_name,
                   class_snapshot AS class_name,
                   status
            FROM session_students
            WHERE session_id=?
            ORDER BY name_snapshot
            """,
            (sid,),
        ).fetchall()
        return {"total": len(rows), "students": [dict(r) for r in rows]}


@teacher.post("/sessions/{sid}/students/{ssid}/reset")
def reset_student(sid: int, ssid: int):
    """Cho làm lại từ đầu (sập nguồn, hỏng máy...): xóa bài làm + kết quả cũ."""
    with db() as c:
        open_session(c, sid)
        cur = c.execute(
            """
            UPDATE session_students 
            SET status='registered', 
                started_at=NULL, 
                submitted_at=NULL 
            WHERE session_id=? 
                AND session_student_id=?
            """,
            (sid, ssid),
        )

        if not cur.rowcount:
            raise HTTPException(
                status_code=404, detail="Không tìm thấy học sinh trong phòng thi"
            )

        c.execute(
            """
            DELETE FROM answers 
            WHERE session_student_id=?
            """,
            (ssid,),
        )

        c.execute(
            """
            DELETE FROM results 
            WHERE session_student_id=?
            """,
            (ssid,),
        )

        return {"status": "ok"}


# phần làm bài của học sinh
class Join(BaseModel):
    access_code: str = Field(min_length=1, max_length=32)
    student_code: str = Field(min_length=1, max_length=50)
    password: str | None = Field(default=None, max_length=128)


def student_ctx(c, ssid):
    r = c.execute(
        """
        SELECT ss.session_student_id, 
        ss.status, ss.started_at, 
        ss.student_code_snapshot,
        ss.name_snapshot,
        ss.class_snapshot,
        ss.session_id, 
        es.version_id, 
        es.paper_code_prefix,
        es.paper_code_length,
        es.paper_variant_count,
        es.end_time, 
        es.status AS session_status,
        e.duration_minutes,
        e.title AS exam_title,
        sub.subject_name,
        v.shuffle_questions, 
        v.shuffle_options
        FROM session_students ss
        JOIN exam_sessions es ON es.session_id = ss.session_id
        JOIN exams e ON e.exam_id = es.exam_id
        JOIN subjects sub ON sub.subject_id = e.subject_id
        JOIN exam_versions v ON v.version_id = es.version_id
        WHERE ss.session_student_id=?
        """,
        (ssid,),
    ).fetchone()

    if not r:
        raise HTTPException(status_code=401, detail="Phiên làm bài không tồn tại")
    return r


def active_ctx(c, ssid):
    """Thí sinh đang được phép làm bài? Hết giờ thì thu bài luôn."""

    r = student_ctx(c, ssid)
    if r["status"] == "submitted":
        raise HTTPException(status_code=403, detail="Bài thi đã được nộp")

    if r["status"] != "in_progress" or r["session_status"] != "ongoing":
        raise HTTPException(
            status_code=403, detail="Bạn không ở trạng thái đang làm bài"
        )

    if now_utc() >= deadline(r["started_at"], r["duration_minutes"], r["end_time"]):
        finalize(c, ssid)
        c.commit()  # phải lưu trước khi raise, vì db() sẽ rollback khi có lỗi
        raise HTTPException(
            status_code=403, detail="Đã hết giờ, bài đã được nộp tự động"
        )
    return r


def current_student(cred=Depends(P.bearer_scheme)):
    try:
        sub = jwt.decode(cred.credentials, P.SECRET_KEY, algorithms=[P.ALGORITHM])[
            "sub"
        ]
        if not sub.startswith("ss:"):
            raise ValueError
        return int(sub[3:])
    except (jwt.InvalidTokenError, KeyError, ValueError):
        raise HTTPException(
            status_code=401,
            detail="Phiên làm bài không hợp lệ, hãy đăng nhập lại",
            headers={"WWW-Authenticate": "Bearer"},
        )


@router.post("/exam/join")
def join(d: Join):
    with db() as c:
        s = c.execute(
            """
            SELECT * 
            FROM exam_sessions 
            WHERE access_code=?
            """,
            (d.access_code.strip().upper(),),
        ).fetchone()

        if not s or s["status"] != "ongoing":
            raise HTTPException(
                status_code=403, detail="Mã phòng không đúng hoặc phòng thi chưa mở"
            )

        now = now_utc()
        if s["start_time"] and now < parse(s["start_time"]):
            raise HTTPException(status_code=403, detail="Chưa đến giờ thi")

        if s["end_time"] and now >= parse(s["end_time"]):
            raise HTTPException(status_code=403, detail="Phòng thi đã đóng")

        st = c.execute(
            """
            SELECT * 
            FROM session_students 
            WHERE session_id=? 
                AND student_code_snapshot=?
            """,
            (s["session_id"], d.student_code.strip()),
        ).fetchone()

        if not st:
            raise HTTPException(
                status_code=404,
                detail="Mã học sinh/SBD không có trong danh sách phòng thi",
            )

        join_deadline = parse(s["start_time"]) + timedelta(minutes=15) if s["start_time"] else None
        if st["status"] == "registered" and join_deadline and now >= join_deadline:
            c.execute(
                "UPDATE session_students SET status='absent' WHERE session_id=? AND status='registered'",
                (s["session_id"],),
            )
            c.commit()
            raise HTTPException(
                status_code=403,
                detail="Đã quá 15 phút từ giờ bắt đầu. Bạn được ghi nhận vắng và không thể vào thi.",
            )
        if st["status"] == "absent":
            raise HTTPException(
                status_code=403,
                detail="Bạn đã được ghi nhận vắng vì vào thi muộn quá 15 phút.",
            )

        if s["password_hash"] and st:
            attempt = c.execute(
                "SELECT failed_attempts, locked_until FROM student_room_login_attempts WHERE session_student_id=?",
                (st["session_student_id"],),
            ).fetchone()
            if attempt and attempt["locked_until"]:
                until = parse(attempt["locked_until"])
                if now < until:
                    minutes = max(1, ceil((until - now).total_seconds() / 60))
                    raise HTTPException(
                        429,
                        f"Đã nhập sai mật khẩu quá 5 lần. Thử lại sau khoảng {minutes} phút.",
                    )
                c.execute(
                    "DELETE FROM student_room_login_attempts WHERE session_student_id=?",
                    (st["session_student_id"],),
                )

        password_ok = not s["password_hash"] or bool(
            d.password and P.password_hash.verify(d.password, s["password_hash"])
        )
        if not password_ok:
            if not st:
                raise HTTPException(403, "Sai mật khẩu phòng thi hoặc mã học sinh/SBD")
            c.execute(
                """
                INSERT INTO student_room_login_attempts (session_student_id, failed_attempts, locked_until)
                VALUES (?, 1, NULL)
                ON CONFLICT(session_student_id) DO UPDATE SET
                    failed_attempts=failed_attempts+1,
                    locked_until=CASE WHEN failed_attempts+1>=5 THEN excluded.locked_until ELSE locked_until END
                """,
                (st["session_student_id"], (now + timedelta(minutes=15)).strftime(FMT)),
            )
            attempt = c.execute(
                "SELECT failed_attempts FROM student_room_login_attempts WHERE session_student_id=?",
                (st["session_student_id"],),
            ).fetchone()
            c.commit()
            if attempt["failed_attempts"] >= 5:
                raise HTTPException(
                    429, "Đã nhập sai mật khẩu 5 lần. Vui lòng thử lại sau 15 phút."
                )
            raise HTTPException(
                403,
                f"Sai mật khẩu phòng thi. Còn {5-attempt['failed_attempts']} lần thử trước khi khóa 15 phút.",
            )

        c.execute(
            "DELETE FROM student_room_login_attempts WHERE session_student_id=?",
            (st["session_student_id"],),
        )

        ssid = st["session_student_id"]
        if st["status"] == "registered":
            c.execute(
                """
                UPDATE session_students 
                SET status='in_progress', 
                    started_at=? 
                WHERE session_student_id=?
                """,
                (P.get_utc_time(), ssid),
            )

        elif (
            st["status"] != "in_progress"
        ):  # in_progress = vào lại sau F5/đứt mạng, được tiếp tục
            raise HTTPException(
                status_code=403,
                detail="Bạn không thể vào thi (đã nộp bài hoặc bị loại)",
            )

        r = active_ctx(c, ssid)
        exp = deadline(
            r["started_at"], r["duration_minutes"], r["end_time"]
        ) + timedelta(minutes=10)
        token = jwt.encode(
            {"sub": f"ss:{ssid}", "exp": exp}, P.SECRET_KEY, algorithm=P.ALGORITHM
        )

        return {
            "access_token": token,
            "token_type": "bearer",
            "full_name": st["name_snapshot"],
            "class_name": st["class_snapshot"],
            "student_code": st["student_code_snapshot"],
        }


def paper_order(c, r):
    ssid = r["session_student_id"]
    ids = [
        x[0]
        for x in c.execute(
            """
        SELECT question_id 
        FROM exam_version_questions 
        WHERE version_id=? 
        ORDER BY question_order
        """,
            (r["version_id"],),
        )
    ]

    if not r["shuffle_questions"]:
        return ids
    if r["paper_variant_count"] <= 1:
        return sorted(ids, key=lambda q: sha(ssid, q))
    variant = (ssid - 1) % r["paper_variant_count"]
    return sorted(ids, key=lambda q: sha("variant", r["session_id"], variant, q))


def paper_code(ssid, settings, shuffled):
    if not shuffled:
        return "GOC"
    if (
        settings["paper_variant_count"] == 1
        and not settings["paper_code_prefix"]
        and settings["paper_code_length"] == 3
    ):
        return str(int(sha("code", ssid)[:6], 16) % 900 + 100)
    variant = (ssid - 1) % settings["paper_variant_count"] + 1
    return settings["paper_code_prefix"] + str(variant).zfill(
        settings["paper_code_length"]
    )


@router.get("/exam/me")
def exam_me(ssid: int = Depends(current_student)):
    """Trạng thái để vào lại bài sau F5: thời gian còn lại + các đáp án đã lưu."""
    with db() as c:
        r = active_ctx(c, ssid)
        question_ids = paper_order(c, r)
        if not question_ids:
            raise HTTPException(
                status_code=409,
                detail="Đề của phòng thi đang trống. Giáo viên cần tạo lại phòng từ đề đã lưu.",
            )
        left = (
            deadline(r["started_at"], r["duration_minutes"], r["end_time"]) - now_utc()
        )
        answers = {
            a["question_id"]: a["answer_text"]
            for a in c.execute(
                """
                       SELECT 
                       question_id, 
                       answer_text 
                       FROM answers 
                       WHERE session_student_id=?
                       """,
                (ssid,),
            )
        }

        return {
            "remaining_seconds": int(left.total_seconds()),
            "paper_code": paper_code(
                ssid, r, r["shuffle_questions"] or r["shuffle_options"]
            ),
            "student_code": r["student_code_snapshot"],
            "full_name": r["name_snapshot"],
            "class_name": r["class_snapshot"],
            "exam_title": r["exam_title"],
            "subject_name": r["subject_name"],
            "question_ids": question_ids,
            "answers": answers,
        }


@router.get("/exam/questions/{n}")
def exam_question(n: int, ssid: int = Depends(current_student)):
    """Tải từng câu theo số thứ tự (1..N): không lộ cả đề và không có đáp án đúng."""

    with db() as c:
        r = active_ctx(c, ssid)
        order = paper_order(c, r)
        if not 1 <= n <= len(order):
            raise HTTPException(status_code=404, detail="Không có câu này")

        qid = order[n - 1]
        q = c.execute(
            """
            SELECT 
            question_id, 
            question_type, 
            question_text 
            FROM questions
            WHERE question_id=?
            """,
            (qid,),
        ).fetchone()

        opts = [
            dict(o)
            for o in c.execute(
                """
            SELECT 
            option_key, 
            option_text 
            FROM question_options 
            WHERE question_id=? 
            ORDER BY sort_order
            """,
                (qid,),
            )
        ]

        if r["shuffle_options"] and q["question_type"] == "multiple_choice":
            if r["paper_variant_count"] > 1:
                variant = (ssid - 1) % r["paper_variant_count"]
                opts.sort(
                    key=lambda o: sha(
                        "variant", r["session_id"], variant, qid, o["option_key"]
                    )
                )
            else:
                opts.sort(key=lambda o: sha(ssid, qid, o["option_key"]))

        media = [
            dict(
                m,
                image_url=(
                    f"/media/{m['file_path']}" if m["media_type"] == "image" else None
                ),
            )
            for m in c.execute(
                """
                    SELECT 
                    media_id, 
                    media_type, 
                    file_path 
                    FROM question_media 
                    WHERE question_id=? 
                    ORDER BY sort_order
                    """,
                (qid,),
            )
        ]

        return {
            "number": n,
            "total": len(order),
            "question": dict(q),
            "options": opts,
            "media": media,
        }


class AnswerSave(BaseModel):
    question_id: int
    answer_text: str  


def check_answer(c, qtype, qid, text):
    if qtype == "multiple_choice":
        ok = c.execute(
            """
            SELECT 1 
            FROM question_options 
            WHERE question_id=? 
                AND option_key=?
            """,
            (qid, text),
        ).fetchone()

    elif qtype == "true_false":
        try:
            g = json.loads(text)
            ok = (
                isinstance(g, dict)
                and set(g) <= {"A", "B", "C", "D"}
                and all(type(v) is bool for v in g.values())
            )
        except json.JSONDecodeError:
            ok = False
    else:
        ok = len(text) <= 500
    if not ok:
        raise HTTPException(status_code=400, detail="Đáp án không hợp lệ")


@router.put("/exam/answers")
def save_answer(d: AnswerSave, ssid: int = Depends(current_student)):
    with db() as c:
        r = active_ctx(c, ssid)
        q = c.execute(
            """
            SELECT 
            q.question_type 
            FROM exam_version_questions evq 
            JOIN questions q ON q.question_id=evq.question_id 
            WHERE evq.version_id=? 
                AND evq.question_id=?
            """,
            (r["version_id"], d.question_id),
        ).fetchone()

        if not q:
            raise HTTPException(status_code=404, detail="Câu hỏi không thuộc đề này")

        text = d.answer_text.strip()
        if not text:
            c.execute(
                """
                DELETE FROM answers 
                WHERE session_student_id=? 
                    AND question_id=?
                """,
                (ssid, d.question_id),
            )

            return {"status": "ok"}

        check_answer(c, q["question_type"], d.question_id, text)
        now = P.get_utc_time()

        c.execute(
            """
            INSERT INTO answers (
            session_student_id, 
            question_id, 
            answer_text, 
            answered_at, 
            updated_at
            ) VALUES (?,?,?,?,?)
            ON CONFLICT(
            session_student_id,
            question_id) 
            DO UPDATE 
            SET answer_text=excluded.answer_text, 
            updated_at=excluded.updated_at
            """,
            (ssid, d.question_id, text, now, now),
        )

        return {"status": "ok"}


@router.post("/exam/submit")
def submit(ssid: int = Depends(current_student)):
    with db() as c:
        if student_ctx(c, ssid)["status"] != "in_progress":
            raise HTTPException(status_code=409, detail="Bài thi đã được nộp")
        finalize(c, ssid)

        return {"status": "ok", "message": "Nộp bài thành công"}


class Violation(BaseModel):
    violation_type: str = Field(
        min_length=1, max_length=50
    )  # ví dụ: tab_hidden, window_blur
    description: str | None = Field(default=None, max_length=300)


@router.post("/exam/violations")
def report_violation(d: Violation, ssid: int = Depends(current_student)):
    with db() as c:
        active_ctx(c, ssid)
        c.execute(
            """
            INSERT INTO exam_violations (
                session_student_id,
                violation_type,
                description
            ) VALUES (?, ?, ?)
            """,
            (ssid, d.violation_type, d.description),
        )

        n = c.execute(
            """
            SELECT COUNT(*)
            FROM exam_violations
            WHERE session_student_id=?
            """,
            (ssid,),
        ).fetchone()[0]

        disqualified = n >= 3
        if disqualified:
            c.execute(
                "UPDATE session_students SET status='disqualified' "
                "WHERE session_student_id=? AND status='in_progress'",
                (ssid,),
            )
            c.execute("DELETE FROM results WHERE session_student_id=?", (ssid,))

        return {"status": "ok", "violation_count": n, "disqualified": disqualified}


# phần chấm điểm 
def score_one(c, qid, qtype, max_score, ans):
    if qtype == "multiple_choice":
        r = c.execute(
            """
            SELECT is_correct 
            FROM question_options 
            WHERE question_id=? 
                AND option_key=?
            """,
            (qid, ans),
        ).fetchone()

        ok = bool(r and r[0])
        return int(ok), max_score if ok else 0.0

    if qtype == "true_false":
        key = {
            k: bool(v)
            for k, v in c.execute(
                """
            SELECT 
            option_key, 
            is_correct 
            FROM question_options 
            WHERE question_id=?
            """,
                (qid,),
            )
        }

        given = json.loads(ans)
        n = sum(
            given.get(k) is v for k, v in key.items()
        )  # đúng 1/2/3/4 ý → 10%/25%/50%/100% điểm câu

        return (
            int(n == 4 and len(key) == 4),
            max_score * (0, 0.1, 0.25, 0.5, 1)[min(n, 4)],
        )

    norm = ans.strip()
    for text, case in c.execute(
        """
        SELECT answer_text, 
        is_case_sensitive 
        FROM question_accepted_answers 
        WHERE question_id=?
        """,
        (qid,),
    ):

        if (
            (norm == text.strip())
            if case
            else (norm.casefold() == text.strip().casefold())
        ):
            return 1, max_score
    return 0, 0.0


def grade(c, ssid):
    """Chấm lại từ đầu (chạy lại nhiều lần cũng ra cùng kết quả)."""

    rows = c.execute(
        """
        SELECT 
        evq.question_id, 
        evq.score, 
        q.question_type, 
        a.answer_text
        FROM session_students ss
        JOIN exam_sessions es ON es.session_id = ss.session_id
        JOIN exam_version_questions evq ON evq.version_id = es.version_id
        JOIN questions q ON q.question_id = evq.question_id
        LEFT JOIN answers a ON a.session_student_id = ss.session_student_id AND a.question_id = evq.question_id
        WHERE ss.session_student_id=?""",
        (ssid,),
    ).fetchall()
    total, right, wrong, blank = 0.0, 0, 0, 0
    for r in rows:
        if r["answer_text"] is None:
            blank += 1
            continue
        ok, sc = score_one(
            c, r["question_id"], r["question_type"], r["score"], r["answer_text"]
        )
        c.execute(
            """
            UPDATE answers 
            SET is_correct=?, 
                score=? 
            WHERE session_student_id=? 
                AND question_id=?
            """,
            (ok, sc, ssid, r["question_id"]),
        )

        total += sc
        right, wrong = right + ok, wrong + (not ok)
    sub = c.execute(
        """
        SELECT submitted_at 
        FROM session_students 
        WHERE session_student_id=?
        """,
        (ssid,),
    ).fetchone()[0]

    c.execute(
        """
        INSERT INTO results (
        session_student_id, 
        total_score, 
        correct_count, 
        incorrect_count, 
        unanswered_count, 
        submitted_at, 
        graded_at
        )
        VALUES (?,?,?,?,?,?,?)
        ON CONFLICT(session_student_id) DO UPDATE SET total_score=excluded.total_score, correct_count=excluded.correct_count,
        incorrect_count=excluded.incorrect_count, unanswered_count=excluded.unanswered_count,
        submitted_at=excluded.submitted_at, graded_at=excluded.graded_at""",
        (ssid, round(total, 2), right, wrong, blank, sub, P.get_utc_time()),
    )


def finalize(c, ssid):
    cur = c.execute(
        """
        UPDATE session_students 
        SET status='submitted', 
            submitted_at=? 
        WHERE session_student_id=? 
            AND status='in_progress'
        """,
        (P.get_utc_time(), ssid),
    )

    if cur.rowcount:
        grade(c, ssid)


def ensure_graded(c, sid):
    """Bài bị vòng auto_submit_loop của python.py thu (chỉ đổi status) sẽ được chấm ở đây."""

    for (ssid,) in c.execute(
        """
        SELECT 
        ss.session_student_id 
        FROM session_students ss
        LEFT JOIN results r ON r.session_student_id = ss.session_student_id
        WHERE ss.session_id=? 
            AND ss.status='submitted'
            AND r.result_id IS NULL
        """,
        (sid,),
    ).fetchall():
        grade(c, ssid)


# phần xuất kết quả
@teacher.get("/sessions/{sid}/results.xlsx")
def export_results(sid: int):
    with db() as c:
        s = get_session(c, sid)
        ensure_graded(c, sid)
        v = c.execute(
            """
            SELECT 
            shuffle_questions, 
            shuffle_options 
            FROM exam_versions 
            WHERE version_id=?
            """,
            (s["version_id"],),
        ).fetchone()

        shuffled = bool(v[0] or v[1])
        rows = c.execute(
            """
            SELECT 
            ss.session_student_id, 
            ss.student_code_snapshot, 
            ss.name_snapshot, 
            ss.class_snapshot, 
            ss.status, 
            ss.submitted_at,
            r.total_score, r.correct_count,

            (
            SELECT COUNT(*) 
            FROM exam_violations x 
            WHERE x.session_student_id = ss.session_student_id
            ) AS violations

            FROM session_students ss LEFT JOIN results r ON r.session_student_id = ss.session_student_id
            WHERE ss.session_id=? ORDER BY ss.session_student_id""",
            (sid,),
        ).fetchall()

        wb = Workbook()
        ws = wb.active
        ws.title = "Bảng điểm"
        ws.append(
            [
                "STT",
                "Mã HS",
                "Họ tên",
                "Lớp",
                "Mã đề",
                "Số câu đúng",
                "Điểm số",
                "Thời gian nộp",
                "Số lần vi phạm",
            ]
        )
        bins = [0] * 10
        for i, r in enumerate(rows, 1):
            done = r["status"] == "submitted"
            when = (
                parse(r["submitted_at"]).astimezone().strftime("%d/%m/%Y %H:%M:%S")
                if r["submitted_at"]
                else ""
            )
            ws.append(
                [
                    i,
                    r[1],
                    r[2],
                    r[3],
                    paper_code(r[0], s, shuffled),
                    r["correct_count"] if done else "",
                    (
                        r["total_score"]
                        if done
                        else ("Vắng" if r["status"] in ("registered", "absent") else "")
                    ),
                    when,
                    r["violations"],
                ]
            )
            if done:
                bins[min(int(r["total_score"]), 9)] += 1

        ws2 = wb.create_sheet("Phổ điểm")
        ws2.append(["Khoảng điểm", "Số học sinh"])
        for i, n in enumerate(bins):
            ws2.append([f"{i} - {i + 1}" if i < 9 else "9 - 10", n])

        submitted = sum(r["status"] == "submitted" for r in rows)
        ws3 = wb.create_sheet("Theo câu hỏi")
        ws3.append(
            [
                "Câu",
                "Mã câu hỏi",
                "Nội dung",
                "Số HS làm đúng",
                "Tỷ lệ đúng",
                "Tỷ lệ sai/bỏ trống",
            ]
        )
        for q in c.execute(
            """
            SELECT 
            evq.question_order, 
            q.question_id, 
            q.question_text, 
            COALESCE(SUM(a.is_correct), 0) AS n_right
            FROM exam_version_questions evq 
            JOIN questions q ON q.question_id = evq.question_id
            LEFT JOIN answers a ON a.question_id = q.question_id
                AND a.session_student_id IN (SELECT session_student_id FROM session_students WHERE session_id=? AND status='submitted')
            WHERE evq.version_id=? 
            GROUP BY evq.question_order 
            ORDER BY evq.question_order
            """,
            (sid, s["version_id"]),
        ):

            rate = q["n_right"] / submitted if submitted else 0
            ws3.append(
                [
                    q["question_order"],
                    q["question_id"],
                    q["question_text"][:120],
                    q["n_right"],
                    round(rate, 3),
                    round(1 - rate, 3) if submitted else 0,
                ]
            )
        return xlsx_response(wb, f"ket_qua_phong_{sid}.xlsx")

router.include_router(teacher)