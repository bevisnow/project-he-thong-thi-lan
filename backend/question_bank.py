import io
import secrets
import sqlite3
from pathlib import Path
from types import SimpleNamespace
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

import python as P
import docx_import
from exam_routes import (
    Difficulty,
    QType,
    cell,
    db,
    validate_question_data,
    xlsx_response,
)

router = APIRouter(prefix="/api", dependencies=[Depends(P.get_current_user)])


class OptionIn(BaseModel):
    option_key: str = Field(min_length=1, max_length=5)
    option_text: str = Field(min_length=1, max_length=5000)
    is_correct: bool = False
    sort_order: int = Field(default=1, gt=0)


class AcceptedIn(BaseModel):
    answer_text: str = Field(min_length=1, max_length=255)
    is_case_sensitive: bool = False


class MediaIn(BaseModel):
    media_type: Literal["image", "file", "audio"]
    file_path: str = Field(min_length=1, max_length=255)  # tên file lấy từ POST /questions/media
    sort_order: int = Field(default=1, gt=0)


class QuestionIn(BaseModel):  # dùng chung cho tạo mới và sửa (PUT thay toàn bộ)
    subject_id: int
    grade_id: int
    topic_id: int
    lesson_id: int | None = None
    question_type: QType
    difficulty: Difficulty
    question_text: str = Field(min_length=1, max_length=20000)
    explanation: str | None = Field(default=None, max_length=20000)
    is_active: bool = True
    is_public: bool = False
    options: list[OptionIn] = Field(default_factory=list, max_length=4)
    accepted_answers: list[AcceptedIn] = Field(default_factory=list, max_length=100)
    media: list[MediaIn] = Field(default_factory=list, max_length=20)


def check_refs(c, d):
    topic = c.execute(
        """
        SELECT subject_id, grade_id
        FROM topics
        WHERE topic_id = ?
            AND is_active = 1
        """,
        (d.topic_id,),
    ).fetchone()

    if not topic or (topic[0], topic[1]) != (d.subject_id, d.grade_id):
        raise HTTPException(
            status_code=400,
            detail="Chủ đề không tồn tại hoặc không thuộc môn/khối đã chọn",
        )

    if d.lesson_id:
        lesson = c.execute(
            """
            SELECT 1
            FROM lessons
            WHERE lesson_id = ?
                AND topic_id = ?
                AND is_active = 1
            """,
            (d.lesson_id, d.topic_id),
        ).fetchone()

        if not lesson:
            raise HTTPException(
                status_code=400,
                detail="Bài học không tồn tại hoặc không thuộc chủ đề đã chọn",
            )

    for media in d.media:
        if (
            Path(media.file_path).name != media.file_path
            or not (P.MEDIA_DIR / media.file_path).is_file()
        ):
            raise HTTPException(
                status_code=400,
                detail=f"File đính kèm '{media.file_path}' chưa được upload",
            )


def used_in_exam(c, qid, only_live=False):
    sql = """
        SELECT 1
        FROM exam_version_questions evq
        WHERE evq.question_id = ?
    """

    if only_live:  # phòng đang/đã thi: đổi đáp án sẽ làm lệch điểm
        sql = """
            SELECT 1
            FROM exam_version_questions evq
            JOIN exam_sessions es
                ON es.version_id = evq.version_id
            WHERE evq.question_id = ?
                AND es.status IN ('ongoing', 'completed')
        """

    return c.execute(sql, (qid,)).fetchone() is not None


def get_owned(c, qid, user):
    question = c.execute(
        """
        SELECT *
        FROM questions
        WHERE question_id = ?
        """,
        (qid,),
    ).fetchone()

    if not question:
        raise HTTPException(
            status_code=404,
            detail="Không tìm thấy câu hỏi",
        )

    if user["role"] != "admin" and question["created_by"] != user["user_id"]:
        raise HTTPException(
            status_code=403,
            detail="Chỉ người tạo câu hỏi hoặc admin mới được sửa/xóa",
        )

    return question


def get_visible(c, qid, user):
    question = c.execute(
        "SELECT * FROM questions WHERE question_id=?", (qid,)
    ).fetchone()
    if not question or (
        user["role"] != "admin"
        and question["created_by"] != user["user_id"]
        and not question["is_public"]
    ):
        raise HTTPException(404, detail="Không tìm thấy câu hỏi")
    return question


def save(c, d, user_id, qid=None):
    validate_question_data(d)
    check_refs(c, d)

    fields = (
        d.subject_id,
        d.grade_id,
        d.topic_id,
        d.lesson_id,
        d.question_type,
        d.difficulty,
        d.question_text.strip(),
        d.explanation,
        int(d.is_active),
        int(d.is_public),
    )

    if qid is None:
        qid = c.execute(
            """
            INSERT INTO questions (
                subject_id,
                grade_id,
                topic_id,
                lesson_id,
                question_type,
                difficulty,
                question_text,
                explanation,
                is_active,
                is_public,
                created_by,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            fields + (user_id, P.get_utc_time()),
        ).lastrowid
    else:
        c.execute(
            """
            UPDATE questions
            SET subject_id = ?,
                grade_id = ?,
                topic_id = ?,
                lesson_id = ?,
                question_type = ?,
                difficulty = ?,
                question_text = ?,
                explanation = ?,
                is_active = ?,
                is_public = ?,
                updated_at = ?
            WHERE question_id = ?
            """,
            fields + (P.get_utc_time(), qid),
        )

        for table in (
            "question_options",
            "question_accepted_answers",
            "question_media",
        ):
            c.execute(
                f"DELETE FROM {table} WHERE question_id = ?",
                (qid,),
            )

    c.executemany(
        """
        INSERT INTO question_options (
            question_id,
            option_key,
            option_text,
            is_correct,
            sort_order
        )
        VALUES (?, ?, ?, ?, ?)
        """,
        [
            (
                qid,
                option.option_key,
                option.option_text,
                int(option.is_correct),
                option.sort_order,
            )
            for option in d.options
        ],
    )

    c.executemany(
        """
        INSERT INTO question_accepted_answers (
            question_id,
            answer_text,
            is_case_sensitive
        )
        VALUES (?, ?, ?)
        """,
        [
            (qid, answer.answer_text, int(answer.is_case_sensitive))
            for answer in d.accepted_answers
        ],
    )

    c.executemany(
        """
        INSERT INTO question_media (
            question_id,
            media_type,
            file_path,
            sort_order
        )
        VALUES (?, ?, ?, ?)
        """,
        [
            (qid, media.media_type, media.file_path, media.sort_order)
            for media in d.media
        ],
    )

    return qid


def conflict(
    func, *args
):  
    try:
        return func(*args)
    except sqlite3.IntegrityError as e:
        raise HTTPException(
            status_code=409,
            detail=f"Dữ liệu trùng hoặc vi phạm ràng buộc: {e}",
        )

@router.get("/questions/stats")
def stats(subject_id: int | None = None, grade_id: int | None = None, user=Depends(P.get_current_user)):
    """Còn bao nhiêu câu đang dùng được theo chủ đề/mức độ/loại: để giáo viên đặt ma trận cho khả thi."""

    with db() as c:

        rows = c.execute(
            """
            SELECT 
            topic_id, 
            lesson_id, 
            difficulty, 
            question_type, 
            COUNT(*) AS n 
            FROM questions
            WHERE is_active=1 AND (? = 'admin' OR created_by = ? OR is_public=1)
                AND (? IS NULL OR subject_id=?) 
                AND (? IS NULL OR grade_id=?)
            GROUP BY topic_id, lesson_id, difficulty, question_type 
            ORDER BY topic_id
            """,
            (user["role"], user["user_id"], subject_id, subject_id, grade_id, grade_id),
        ).fetchall()

        return {"stats": [dict(r) for r in rows]}


@router.post("/questions/media", status_code=201)
def upload_media(file: UploadFile = File(...)):
    ext = Path(file.filename or "").suffix.lower()
    if ext not in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        raise HTTPException(
            status_code=400,
            detail="Chỉ nhận ảnh png, jpg, jpeg, gif, webp",
        )
    data = P.read_upload_limited(file, 5 * 1024 * 1024, "Ảnh tối đa 5MB")
    try:
        data = docx_import.sanitize_image(data, ext)
    except RuntimeError as exc:
        raise HTTPException(503, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    name = secrets.token_hex(8) + ext
    (P.MEDIA_DIR / name).write_bytes(data)
    return {"file_path": name, "image_url": f"/media/{name}"}


KIND = {"TN": "multiple_choice", "DS": "true_false", "TLN": "short_answer"}


@router.get("/templates/questions.xlsx")
def questions_template():
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Loại (TN/DS/TLN)",
            "Độ khó (NB/TH/VD/VDC)",
            "Nội dung",
            "A",
            "B",
            "C",
            "D",
            "Đáp án",
            "Giải thích",
        ]
    )
    ws.append(["TN", "NB", "1 + 1 = ?", "1", "2", "3", "4", "B", ""])
    ws.append(
        [
            "DS",
            "TH",
            "Cho hàm số y = x². Các mệnh đề sau đúng hay sai?",
            "Đồ thị là parabol",
            "Hàm chẵn",
            "Có cực đại",
            "Đi qua gốc tọa độ",
            "A,B,D",
            "",
        ]
    )
    ws.append(
        ["TLN", "VD", "Thủ đô của Việt Nam?", "", "", "", "", "Hà Nội;Ha Noi", ""]
    )
    return xlsx_response(wb, "cau_hoi_mau.xlsx")


def row_to_question(ctx, r):
    kind, diff, text, a, b, c_, d_, ans, expl = (
        cell(x) for x in (list(r) + [None] * 9)[:9]
    )
    qtype = KIND.get(kind.upper())
    if not qtype or diff.upper() not in ("NB", "TH", "VD", "VDC"):
        raise ValueError("Loại phải là TN/DS/TLN và độ khó phải là NB/TH/VD/VDC")
    opts, accepted = [], []
    if qtype == "short_answer":
        accepted = [
            AcceptedIn(answer_text=x.strip()) for x in ans.split(";") if x.strip()
        ]
    else:
        right = {
            x.strip().upper() for x in ans.replace(";", ",").split(",") if x.strip()
        }
        opts = [
            OptionIn(option_key=k, option_text=t, is_correct=k in right, sort_order=i)
            for i, (k, t) in enumerate(zip("ABCD", (a, b, c_, d_)), 1)
            if t
        ]
        if right - {o.option_key for o in opts}:
            raise ValueError("Cột Đáp án trỏ tới lựa chọn bị bỏ trống")
    d = QuestionIn(
        **ctx,
        question_type=qtype,
        difficulty=diff.upper(),
        question_text=text,
        explanation=expl or None,
        options=opts,
        accepted_answers=accepted,
    )
    validate_question_data(d)
    return d


def docx_to_question(ctx, r):
    qtype, diff, stem, opts, accepted = docx_import.parse_row(r)
    d = QuestionIn(
        **ctx,
        question_type=qtype,
        difficulty=diff,
        question_text=stem,
        options=[
            OptionIn(option_key=k, option_text=t, is_correct=ok, sort_order=i)
            for i, (k, t, ok) in enumerate(opts, 1)
        ],
        accepted_answers=[AcceptedIn(answer_text=x) for x in accepted],
    )
    validate_question_data(d)
    return d


@router.get("/templates/questions-{kind}.docx")
def docx_template(kind: str):
    """Tải file Word mẫu (tn / ds / tln) trong thư mục templates/ cạnh python.py."""
    if kind not in ("tn", "ds", "tln"):
        raise HTTPException(
            status_code=404,
            detail="Không có mẫu này",
        )
    f = P.BASE_DIR / "templates" / f"Loai_{kind.upper()}.docx"
    if not f.is_file():
        raise HTTPException(
            status_code=404,
            detail="Chưa đặt file mẫu vào thư mục templates",
        )
    return Response(
        f.read_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{f.name}"'},
    )


@router.post("/questions/import")
def import_questions(
    subject_id: int,
    grade_id: int,
    topic_id: int,
    lesson_id: int | None = None,
    is_public: bool = False,
    commit: bool = False,
    file: UploadFile = File(...),
    user: dict = Depends(P.get_current_user),
):
    """Nhận .xlsx hoặc .docx (mẫu Loai_TN/DS/TLN). Môn/khối/chủ đề/bài chọn trên web cho cả file.
    commit=false: xem trước + lỗi từng dòng; commit=true: ghi, chỉ khi không còn lỗi."""
    ctx = dict(
        subject_id=subject_id, grade_id=grade_id, topic_id=topic_id, lesson_id=lesson_id, is_public=is_public
    )
    raw = P.read_upload_limited(file, 10 * 1024 * 1024, "File import tối đa 10MB")
    items = []  # (số dòng, đoạn xem trước, hàm dựng câu hỏi, dữ liệu dòng)
    doc = None
    if (file.filename or "").lower().endswith(".docx"):
        try:
            doc = docx_import.Doc(raw, P.MEDIA_DIR if commit else None)
            for r in doc.rows():
                items.append(
                    (r["row"], (r.get("text") or r["code"])[:100], docx_to_question, r)
                )
        except RuntimeError as exc:
            raise HTTPException(status_code=503, detail=str(exc)) from exc
        except Exception:
            raise HTTPException(
                status_code=400,
                detail="File không phải Word (.docx) đúng mẫu",
            )
    else:
        try:
            docx_import.validate_package(raw)
            ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        except Exception:
            raise HTTPException(
                status_code=400,
                detail="File phải là Excel (.xlsx) hoặc Word (.docx)",
            )
        for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if i > 5001:
                raise HTTPException(413, "Mỗi lần import tối đa 5.000 câu hỏi")
            if any(cell(x) for x in r):
                items.append((i, cell(r[2])[:100], row_to_question, r))
    rows, good = [], []
    for i, preview, build, data in items:
        try:
            good.append(build(ctx, data))
            rows.append({"row": i, "text": preview, "error": None})
        except (ValueError, HTTPException) as e:
            rows.append(
                {"row": i, "text": preview, "error": str(getattr(e, "detail", e))}
            )
    errors = [r for r in rows if r["error"]]
    with db() as c:
        check_refs(c, SimpleNamespace(**ctx, media=[]))
        if not commit:
            return {"count": len(rows), "error_count": len(errors), "rows": rows}
        if errors or not good:
            raise HTTPException(
                400,
                {
                    "message": "File còn lỗi hoặc rỗng, hãy xem trước và sửa lại",
                    "errors": errors,
                },
            )
        for d in good:
            conflict(save, c, d, user["user_id"])
    if doc:
        doc.save_images()
    return {"status": "ok", "imported": len(good)}


@router.get("/questions")
def list_questions(
    subject_id: int | None = None,
    grade_id: int | None = None,
    topic_id: int | None = None,
    lesson_id: int | None = None,
    difficulty: Difficulty | None = None,
    question_type: QType | None = None,
    is_active: int | None = None,
    q: str | None = Query(default=None, max_length=200),
    limit: int = 50,
    offset: int = 0,
    user=Depends(P.get_current_user),
):

    where, args = [], []
    if user["role"] != "admin":
        where.append("(created_by=? OR is_public=1)")
        args.append(user["user_id"])
    for col, v in (
        ("subject_id", subject_id),
        ("grade_id", grade_id),
        ("topic_id", topic_id),
        ("lesson_id", lesson_id),
        ("difficulty", difficulty),
        ("question_type", question_type),
        ("is_active", is_active),
    ):

        if v is not None:
            where.append(f"{col}=?")
            args.append(v)

    if q:
        where.append("question_text LIKE ?")
        args.append(f"%{q}%")
    w = (" WHERE " + " AND ".join(where)) if where else ""

    with db() as c:
        total = c.execute(f"SELECT COUNT(*) FROM questions{w}", args).fetchone()[0]
        rows = c.execute(
            f"SELECT question_id, subject_id, grade_id, topic_id, lesson_id, question_type, difficulty, question_text, "
            f"is_active, is_public, created_by FROM questions{w} ORDER BY question_id DESC LIMIT ? OFFSET ?",
            args + [max(1, min(limit, 200)), max(0, offset)],
        ).fetchall()
        questions = []
        for row in rows:
            item = dict(row)
            item["can_manage"] = user["role"] == "admin" or item["created_by"] == user["user_id"]
            item.pop("created_by")
            questions.append(item)
        return {"total": total, "questions": questions}


@router.get("/questions/{question_id}")
def question_detail(question_id: int, user=Depends(P.get_current_user)):
    """Bản cho giáo viên: có đáp án đúng. Học sinh KHÔNG dùng route này (dùng /api/exam/questions/{n})."""
    with db() as c:
        q = get_visible(c, question_id, user)
        can_manage = user["role"] == "admin" or q["created_by"] == user["user_id"]
        rows = lambda sql: [dict(r) for r in c.execute(sql, (question_id,))]
        media = rows(
            "SELECT media_id, media_type, file_path, sort_order FROM question_media WHERE question_id=? ORDER BY sort_order"
        )
        for m in media:
            m["image_url"] = (
                f"/media/{m['file_path']}" if m["media_type"] == "image" else None
            )

        return {
            "question": dict(q),
            "can_manage": can_manage,
            "used_in_exam": used_in_exam(c, question_id),
            "media": media,
            "options": rows(
                "SELECT option_key, option_text, is_correct, sort_order FROM question_options WHERE question_id=? ORDER BY sort_order"
            ),
            "accepted_answers": rows(
                "SELECT answer_text, is_case_sensitive FROM question_accepted_answers WHERE question_id=?"
            ),
        }


@router.post("/questions", status_code=201)
def create_question(d: QuestionIn, user: dict = Depends(P.get_current_user)):
    with db() as c:
        return {"question_id": conflict(save, c, d, user["user_id"])}


@router.put("/questions/{question_id}")
def update_question(
    question_id: int, d: QuestionIn, user: dict = Depends(P.get_current_user)
):
    with db() as c:
        get_owned(c, question_id, user)
        if used_in_exam(c, question_id, only_live=True):
            raise HTTPException(
                status_code=409,
                detail="Câu hỏi đang/đã được dùng trong phòng thi nên không sửa được, chỉ có thể tắt (is_active=0)",
            )
        conflict(save, c, d, user["user_id"], question_id)
        return {"status": "ok", "question_id": question_id}


class Visibility(BaseModel):
    is_public: bool


@router.patch("/questions/{question_id}/visibility")
def set_visibility(question_id: int, d: Visibility, user: dict = Depends(P.get_current_user)):
    with db() as c:
        get_owned(c, question_id, user)
        c.execute(
            "UPDATE questions SET is_public=?, updated_at=? WHERE question_id=?",
            (int(d.is_public), P.get_utc_time(), question_id),
        )
        return {"status": "ok", "is_public": d.is_public}


class Active(BaseModel):
    is_active: bool


@router.patch("/questions/{question_id}/active")
def set_active(question_id: int, d: Active, user: dict = Depends(P.get_current_user)):
    with db() as c:
        get_owned(c, question_id, user)
        c.execute(
            "UPDATE questions SET is_active=?, updated_at=? WHERE question_id=?",
            (int(d.is_active), P.get_utc_time(), question_id),
        )
        return {"status": "ok", "is_active": d.is_active}


@router.delete("/questions/{question_id}")
def delete_question(question_id: int, user: dict = Depends(P.get_current_user)):
    """Chưa dùng trong đề nào → xóa hẳn; đã dùng → chỉ tắt để giữ nguyên đề và kết quả cũ."""
    with db() as c:
        get_owned(c, question_id, user)
        if used_in_exam(c, question_id):
            c.execute(
                "UPDATE questions SET is_active=0, updated_at=? WHERE question_id=?",
                (P.get_utc_time(), question_id),
            )
            return {"status": "ok", "deleted": False, "deactivated": True}
        c.execute("DELETE FROM questions WHERE question_id=?", (question_id,))
        return {"status": "ok", "deleted": True}
